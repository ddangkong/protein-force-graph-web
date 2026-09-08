# -*- coding: utf-8 -*-
"""Build the Excel workbook from out/master.json (run analysis/build.py first).
Summary sheets use COUNTIFS/SUMIFS formulas over the data sheet (wildcard substring criteria),
plus a static '수집 시점 값' cross-check column so numbers show even before recalculation."""
import json, os, sys, datetime, collections
import re as _re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from tags import TAG_RULES, CATEGORY_LABEL, tag_text
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.workbook.properties import CalcProperties

_ZW = _re.compile("[​-‏  ﻿]")


def clean(v):
    """Strip characters Excel cannot store (control chars, zero-width marks); neutralise formula-like text."""
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
        v = _ZW.sub("", v)
        if v and v[0] in "=+-@":
            v = "'" + v
    return v


SCRATCH = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(SCRATCH, "out")
items = json.load(open(os.path.join(OUT, "master.json"), encoding="utf-8"))
items.sort(key=lambda x: (x.get("date") or "0000-00-00"), reverse=True)
SUBSET = os.environ.get("XLSX_SUBSET")  # verification mode: mixed slice, written to _reduced.xlsx
if SUBSET:
    yt = [i for i in items if i["platform"] == "YouTube" and (i.get("date") or "")[:4] in ("2025", "2026")][:300]
    ot = [i for i in items if i["platform"] != "YouTube"][:250]
    items = yt + ot

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)
BOLD = Font(name=FONT, bold=True, size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=14)
NOTE_FONT = Font(name=FONT, italic=True, size=9, color="555555")
STATIC_FONT = Font(name=FONT, size=10, color="0000FF")  # blue = hardcoded (수집 시점 값)
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

DATA_COLS = [
    ("플랫폼", 11), ("게시일", 11), ("형식", 9), ("제목/캡션", 60), ("작성자(채널/계정)", 24), ("조회수", 11), ("좋아요", 9),
    ("댓글수", 8), ("튜닝 태그", 40), ("튜닝 카테고리", 30), ("짜파게티 언급", 8), ("레시피성", 8), ("먹방", 7),
    ("URL", 45), ("설명/레시피 본문(발췌)", 80), ("수집 검색어", 30), ("상대시점/기타", 14),
]
DATA_SHEET = "전체_최신순"
LAST = len(items) + 1  # last data row
D = f"'{DATA_SHEET}'"
PLATS = ["YouTube", "TikTok", "Instagram", "Web"]


def rng(col):
    return f"{D}!${col}$2:${col}${LAST}"


def cf(*pairs):
    """COUNTIFS over pairs of (col, criteria)."""
    parts = []
    for col, crit in pairs:
        parts.append(f'{rng(col)},"{crit}"')
    return "=COUNTIFS(" + ",".join(parts) + ")"


def sf(sum_col, *pairs):
    parts = [rng(sum_col)]
    for col, crit in pairs:
        parts.append(f'{rng(col)},"{crit}"')
    return "=SUMIFS(" + ",".join(parts) + ")"


# ---- Python-side reference counts
def _yr(it):
    return (it.get("date") or "")[:4]


def has_tag(it, tag):
    return tag in "; ".join(it.get("tags") or [])


def has_cat(it, cat):
    return cat in "; ".join(CATEGORY_LABEL.get(c, c) for c in it.get("categories") or [])


def py_count(pred):
    return sum(1 for it in items if pred(it))


def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER


def fmt_block(ws, r0, r1, c0, c1, static_cols=()):
    for rr in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            cell = ws.cell(row=rr, column=c)
            cell.border = BORDER
            if c > 1 and cell.font != BOLD:
                cell.font = STATIC_FONT if c in static_cols else BODY_FONT
            if c > 1:
                cell.number_format = "#,##0"


def write_data_sheet(ws, rows):
    for i, (name, width) in enumerate(DATA_COLS, 1):
        ws.cell(row=1, column=i, value=name)
        ws.column_dimensions[get_column_letter(i)].width = width
    style_header(ws, len(DATA_COLS))
    for r, it in enumerate(rows, 2):
        vals = [
            it["platform"], it.get("date") or "", it.get("format") or "", (it.get("title") or "")[:1500],
            it.get("author") or "", it.get("views"), it.get("likes"), it.get("comments_count"),
            "; ".join(it.get("tags") or []), "; ".join(CATEGORY_LABEL.get(c, c) for c in it.get("categories") or []),
            "Y" if it.get("mentions_jjapagetti") else "N", "Y" if it.get("recipe_like") else "N",
            "Y" if it.get("is_mukbang") else "N", it.get("url"), (it.get("description") or "")[:3000],
            "; ".join(it.get("queries") or [])[:500], (it.get("relative_date") or "")[:200],
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=clean(v))
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (4, 9, 10, 15)))
        ws.cell(row=r, column=14).hyperlink = clean(it.get("url"))
        ws.cell(row=r, column=14).font = Font(name=FONT, size=10, color="0563C1", underline="single")
        for c in (6, 7, 8):
            ws.cell(row=r, column=c).number_format = "#,##0"
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DATA_COLS))}{len(rows) + 1}"


def build():
    wb = Workbook()
    # ---------- 요약
    ws = wb.active
    ws.title = "요약"
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFGHIJK":
        ws.column_dimensions[col].width = 14
    ws["A1"] = "짜파게티 SNS 튜닝 레시피 수집 요약"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"수집 기준일: {datetime.date.today().isoformat()} / 대상: YouTube·TikTok·Instagram·웹(기사·커뮤니티·레시피 사이트)"
    ws["A2"].font = NOTE_FONT
    ws["A3"] = "검은색 수치 = '전체_최신순' 시트를 참조하는 COUNTIFS/SUMIFS 수식(파일을 열면 자동 계산). 파란색 수치 = 수집 시점에 파이썬으로 집계한 검증값(고정)."
    ws["A3"].font = NOTE_FONT
    r = 5
    ws.cell(row=r, column=1, value="플랫폼별 수집 건수").font = BOLD
    heads = ["플랫폼", "전체 수집", "짜파게티 언급", "레시피성(튜닝 정보 有)", "2026년 게시", "2025년 게시", "먹방 콘텐츠", "전체 수집(검증값)", "짜파게티 언급(검증값)"]
    for c, h in enumerate(heads, 1):
        ws.cell(row=r + 1, column=c, value=h)
    style_header(ws, len(heads), row=r + 1)
    for i, p in enumerate(PLATS):
        rr = r + 2 + i
        ws.cell(row=rr, column=1, value=p)
        ws.cell(row=rr, column=2, value=cf(("A", p)))
        ws.cell(row=rr, column=3, value=cf(("A", p), ("K", "Y")))
        ws.cell(row=rr, column=4, value=cf(("A", p), ("K", "Y"), ("L", "Y")))
        ws.cell(row=rr, column=5, value=cf(("A", p), ("K", "Y"), ("B", "2026*")))
        ws.cell(row=rr, column=6, value=cf(("A", p), ("K", "Y"), ("B", "2025*")))
        ws.cell(row=rr, column=7, value=cf(("A", p), ("M", "Y")))
        ws.cell(row=rr, column=8, value=py_count(lambda it, p=p: it["platform"] == p))
        ws.cell(row=rr, column=9, value=py_count(lambda it, p=p: it["platform"] == p and it["mentions_jjapagetti"]))
    tr = r + 2 + len(PLATS)
    ws.cell(row=tr, column=1, value="합계")
    for c in range(2, 10):
        L = get_column_letter(c)
        ws.cell(row=tr, column=c, value=f"=SUM({L}{r + 2}:{L}{tr - 1})")
    fmt_block(ws, r + 2, tr, 1, 9, static_cols=(8, 9))
    for c in range(1, 10):
        ws.cell(row=tr, column=c).font = BOLD
    # 카테고리 요약
    r = tr + 2
    ws.cell(row=r, column=1, value="튜닝 카테고리별 건수 (짜파게티 언급 게시물, 중복 태깅 허용)").font = BOLD
    heads = ["카테고리", "전체", "2026", "2025", "2024 이전", "YouTube", "TikTok", "Instagram", "Web", "전체(검증값)"]
    for c, h in enumerate(heads, 1):
        ws.cell(row=r + 1, column=c, value=h)
    style_header(ws, len(heads), row=r + 1)
    cats = list(CATEGORY_LABEL.values())
    for i, cat in enumerate(cats):
        rr = r + 2 + i
        ws.cell(row=rr, column=1, value=cat)
        crit = f"*{cat}*"
        ws.cell(row=rr, column=2, value=cf(("J", crit), ("K", "Y")))
        ws.cell(row=rr, column=3, value=cf(("J", crit), ("K", "Y"), ("B", "2026*")))
        ws.cell(row=rr, column=4, value=cf(("J", crit), ("K", "Y"), ("B", "2025*")))
        ws.cell(row=rr, column=5, value=f"=B{rr}-C{rr}-D{rr}-" + cf(("J", crit), ("K", "Y"), ("B", ""))[1:])
        for j, p in enumerate(PLATS):
            ws.cell(row=rr, column=6 + j, value=cf(("J", crit), ("K", "Y"), ("A", p)))
        ws.cell(row=rr, column=10, value=py_count(lambda it, cat=cat: it["mentions_jjapagetti"] and has_cat(it, cat)))
    fmt_block(ws, r + 2, r + 1 + len(cats), 1, 10, static_cols=(10,))
    r = r + 2 + len(cats) + 1
    ws.cell(row=r, column=1, value="※ 태그/카테고리는 제목·캡션·설명·요약 텍스트에 대한 키워드 규칙 기반 자동 분류이며, 한 게시물에 여러 태그가 붙을 수 있습니다. '2024 이전' = 전체 − 2026 − 2025 − 게시일 미확인. 세부 태그 통계는 '튜닝유형_통계' 시트 참조.").font = NOTE_FONT

    # ---------- 전체_최신순
    ws_all = wb.create_sheet(DATA_SHEET)
    write_data_sheet(ws_all, items)
    # ---------- 플랫폼별
    for p in PLATS:
        wsp = wb.create_sheet(p if p != "Web" else "Web_기사·커뮤니티")
        write_data_sheet(wsp, [i for i in items if i["platform"] == p])

    # ---------- 튜닝유형_통계
    wst = wb.create_sheet("튜닝유형_통계")
    heads = ["카테고리", "튜닝 태그", "전체", "2026", "2025", "2024 이전", "YouTube", "TikTok", "Instagram", "Web", "YouTube 조회수 합(2025~)", "전체(검증값)", "댓글 언급 수(검증값)"]
    for c, h in enumerate(heads, 1):
        wst.cell(row=1, column=c, value=h)
    style_header(wst, len(heads))
    wst.column_dimensions["A"].width = 26
    wst.column_dimensions["B"].width = 34
    for col in "CDEFGHIJ":
        wst.column_dimensions[col].width = 11
    wst.column_dimensions["K"].width = 20
    wst.column_dimensions["L"].width = 14
    wst.column_dimensions["M"].width = 16
    comment_tag_counts = collections.Counter()
    for it in items:
        for cm in it.get("comments") or []:
            for t, _ in tag_text(cm.get("text") or ""):
                comment_tag_counts[t] += 1
    for i, (tag, cat, _) in enumerate(TAG_RULES):
        rr = i + 2
        wst.cell(row=rr, column=1, value=CATEGORY_LABEL.get(cat, cat))
        wst.cell(row=rr, column=2, value=tag)
        crit = f"*{tag}*"
        wst.cell(row=rr, column=3, value=cf(("I", crit), ("K", "Y")))
        wst.cell(row=rr, column=4, value=cf(("I", crit), ("K", "Y"), ("B", "2026*")))
        wst.cell(row=rr, column=5, value=cf(("I", crit), ("K", "Y"), ("B", "2025*")))
        wst.cell(row=rr, column=6, value=f"=C{rr}-D{rr}-E{rr}-" + cf(("I", crit), ("K", "Y"), ("B", ""))[1:])
        for j, p in enumerate(PLATS):
            wst.cell(row=rr, column=7 + j, value=cf(("I", crit), ("K", "Y"), ("A", p)))
        wst.cell(row=rr, column=11, value=sf("F", ("I", crit), ("K", "Y"), ("A", "YouTube"), ("B", "2025*"))[:-1] + ")+" +
                 sf("F", ("I", crit), ("K", "Y"), ("A", "YouTube"), ("B", "2026*"))[1:])
        wst.cell(row=rr, column=12, value=py_count(lambda it, tag=tag: it["mentions_jjapagetti"] and has_tag(it, tag)))
        wst.cell(row=rr, column=13, value=comment_tag_counts.get(tag, 0))
    fmt_block(wst, 2, len(TAG_RULES) + 1, 1, 13, static_cols=(12, 13))
    for rr in range(2, len(TAG_RULES) + 2):
        wst.cell(row=rr, column=1).font = BODY_FONT
        wst.cell(row=rr, column=2).font = BODY_FONT
    wst.freeze_panes = "C2"
    wst.auto_filter.ref = f"A1:M{len(TAG_RULES) + 1}"

    # ---------- 월별 추이
    wsm = wb.create_sheet("월별_추이")
    heads = ["월", "전체", "YouTube", "TikTok", "Instagram", "Web", "전체(검증값)"]
    for c, h in enumerate(heads, 1):
        wsm.cell(row=1, column=c, value=h)
    style_header(wsm, len(heads))
    months = []
    y, m = 2024, 1
    today = datetime.date.today()
    while (y, m) <= (today.year, today.month):
        months.append(f"{y}-{m:02d}")
        m += 1
        if m > 12:
            y, m = y + 1, 1
    for i, mo in enumerate(months):
        rr = i + 2
        wsm.cell(row=rr, column=1, value=mo)
        wsm.cell(row=rr, column=2, value=cf(("B", f"{mo}*"), ("K", "Y")))
        for j, p in enumerate(PLATS):
            wsm.cell(row=rr, column=3 + j, value=cf(("B", f"{mo}*"), ("K", "Y"), ("A", p)))
        wsm.cell(row=rr, column=7, value=py_count(lambda it, mo=mo: it["mentions_jjapagetti"] and (it.get("date") or "").startswith(mo)))
    fmt_block(wsm, 2, len(months) + 1, 1, 7, static_cols=(7,))
    for rr in range(2, len(months) + 2):
        wsm.cell(row=rr, column=1).font = BODY_FONT
    wsm.cell(row=len(months) + 3, column=1, value="※ 게시일이 확인된 짜파게티 언급 게시물 기준. YouTube는 검색 필터(오늘/이번 주/이번 달/올해) 특성상 최근 12개월이 상대적으로 촘촘히 수집됨. TikTok·Instagram은 검색 인덱스 노출분만 포함.").font = NOTE_FONT
    wsm.column_dimensions["A"].width = 12
    wsm.column_dimensions["G"].width = 14

    # ---------- YouTube 댓글
    wsc = wb.create_sheet("YouTube_댓글")
    heads = ["videoId", "영상 제목", "영상 게시일", "댓글 시점", "작성자", "좋아요", "답글수", "댓글 내용", "튜닝 태그"]
    widths = [14, 40, 11, 11, 18, 8, 7, 80, 30]
    for c, (h, w) in enumerate(zip(heads, widths), 1):
        wsc.cell(row=1, column=c, value=h)
        wsc.column_dimensions[get_column_letter(c)].width = w
    style_header(wsc, len(heads))
    rr = 2
    for it in items:
        for cm in it.get("comments") or []:
            vals = [it["id"], (it.get("title") or "")[:200], it.get("date"), cm.get("publishedTime"), cm.get("author"),
                    cm.get("likes"), cm.get("replyCount"), (cm.get("text") or "")[:2000],
                    "; ".join(t for t, _ in tag_text(cm.get("text") or ""))]
            for c, v in enumerate(vals, 1):
                cell = wsc.cell(row=rr, column=c, value=clean(v))
                cell.font = BODY_FONT
                cell.alignment = Alignment(vertical="top", wrap_text=(c == 8))
            rr += 1
    wsc.freeze_panes = "A2"
    wsc.auto_filter.ref = f"A1:I{max(rr - 1, 1)}"

    # ---------- 방법론
    wsn = wb.create_sheet("수집방법·유의사항")
    wsn.column_dimensions["A"].width = 120
    notes = [
        ("수집 방법", TITLE_FONT),
        ("• YouTube: YouTube 공식 InnerTube API(youtubei.googleapis.com) 검색 엔드포인트로 71개 키워드 × 6개 필터(오늘/이번 주/이번 달/올해/조회수순/관련도순) 조합 검색 후 영상별 상세(제목·채널·게시일·조회수·좋아요·설명·댓글) 수집.", BODY_FONT),
        ("• TikTok / Instagram: 세션 네트워크 정책상 tiktok.com·instagram.com 직접 접근이 차단되어, 웹 검색 엔진 인덱스(site:tiktok.com / site:instagram.com 키워드 검색)에 노출된 게시물의 URL·캡션·요약을 수집. 조회수·좋아요는 검색 결과에 노출된 경우에만 기록.", BODY_FONT),
        ("• TikTok 게시일은 영상 ID(상위 32비트 = Unix timestamp), Instagram 게시일은 숏코드(미디어 ID 내장 타임스탬프)에서 복원한 값(UTC 기준, ±1일 오차 가능).", BODY_FONT),
        ("• Web: 뉴스·커뮤니티·레시피 사이트·공식 채널을 웹 검색으로 수집(요리대회 우승 레시피, 신제품, 셰프 레시피, 커뮤니티 꿀조합 등).", BODY_FONT),
        ("• 튜닝 태그/카테고리: 제목·캡션·설명·요약 텍스트에 대한 키워드 규칙(정규식) 기반 자동 분류. 오탐/누락 가능 → 원문(URL) 확인 권장.", BODY_FONT),
        ("• 요약·통계 시트의 검은색 수치는 COUNTIFS/SUMIFS 수식(와일드카드 부분일치)이며 파일을 열 때 계산됩니다. 파란색 수치는 수집 시점의 파이썬 집계값(고정)으로 수식 검증용입니다.", BODY_FONT),
        ("유의사항", TITLE_FONT),
        ("• 검색 엔진 인덱스에 잡히지 않은 TikTok/Instagram 게시물은 포함되지 않음(특히 최근 1~2주 게시물은 인덱싱 지연 가능).", BODY_FONT),
        ("• YouTube 조회수·좋아요·댓글수는 수집 시점 기준 스냅숏.", BODY_FONT),
        ("• '먹방' 플래그는 제목·설명에 먹방/mukbang/ASMR 키워드가 있는 경우로, 레시피 정보가 함께 있을 수 있음.", BODY_FONT),
        (f"• 수집 기준일: {datetime.date.today().isoformat()}", BODY_FONT),
    ]
    for i, (t, f) in enumerate(notes, 1):
        cell = wsn.cell(row=i, column=1, value=t)
        cell.font = f
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    path = os.path.join(OUT, "_reduced.xlsx" if SUBSET else "짜파게티_SNS_튜닝레시피_수집.xlsx")
    wb.save(path)
    print("saved", path, "items", len(items))
    return path


if __name__ == "__main__":
    build()
